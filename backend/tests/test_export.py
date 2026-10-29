"""Exports (docs/API.md §5): columns, CSV and XLSX safety, limits and the export slot."""

from __future__ import annotations

import csv
import io
from collections.abc import Iterator
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook

from mdt.domain import export
from mdt.domain.errors import NotFoundError, ValidationError
from mdt.domain.export import (
    XLSX_MAX_CHARS,
    csv_cell,
    export_filename,
    start_export,
    xlsx_text,
)
from mdt.domain.registry import ENTITY_TYPES
from mdt.domain.revisions import apply_change, soft_delete
from tests.conftest import Db
from tests.helpers import USER, create_event, create_project, create_service
from tests.read_helpers import (
    assert_consistent,
    count_statements,
    create,
    create_document,
    open_copy,
)

LONG_NOTE = "x" * 40_000


@dataclass
class Catalog:
    db: Db
    ids: dict[str, str]

    def csv_rows(self, entity_type: str, **params: str) -> list[list[str]]:
        raw = self.csv_bytes(entity_type, **params)
        assert raw.startswith(b"\xef\xbb\xbf")
        return list(csv.reader(io.StringIO(raw.decode("utf-8-sig"), newline="")))

    def csv_bytes(self, entity_type: str, **params: str) -> bytes:
        with self.db.session() as session:
            job = start_export(session, entity_type, {"format": "csv", **params})
            return b"".join(job.iter_csv())

    def xlsx(self, entity_type: str, **params: str) -> Any:
        buffer = io.BytesIO()
        with self.db.session() as session:
            job = start_export(session, entity_type, {"format": "xlsx", **params})
            job.write_xlsx(buffer)
        buffer.seek(0)
        return load_workbook(buffer)


def build(session: Any) -> dict[str, str]:
    ids: dict[str, str] = {}
    ids["X1"] = create_project(
        session,
        slug="exp-alpha",
        name_en='=HYPERLINK("https://example.org")',
        name_uk="Альфа",
        status="active",
        policy_domains=["domain_b", "domain_a"],
        official_urls=["https://example.org/one", "https://example.org/two"],
        aliases=[{"text": "+Alias two"}, {"text": "Alias one"}],
        in_paper_scope=True,
        scope_note="Synthetic scope note.",
    ).entity_id
    ids["E1"] = create_event(session, ids["X1"], "launch", "2012-05").entity_id
    ids["X2"] = create_project(session, slug="exp-beta", name_en="-5", name_uk="@mention").entity_id
    ids["X3"] = create_project(session, slug="exp-gamma", name_en="#N/A").entity_id
    gone = create_project(session, slug="exp-deleted", name_en="Deleted export")
    soft_delete(session, "project", gone.entity_id, USER)
    ids["X4"] = gone.entity_id
    ids["S1"] = create_service(session, ids["X1"], name_en="Exported service").entity_id
    ids["I1"] = create(
        session, "indicator", subject_id=ids["X2"], metric="Budget", value=-2.5
    ).entity_id
    ids["D1"] = create_document(
        session,
        "Exported document",
        links=[{"subject_id": ids["X2"], "state": "confirmed"}],
    ).entity_id
    ids["T1"] = apply_change(
        session,
        "vocabulary_term",
        None,
        {
            "vocabulary": "status",
            "code": "exp_long",
            "label_en": "Long attributes",
            "attributes": {"note": LONG_NOTE},
        },
        USER,
    ).entity_id
    return ids


@pytest.fixture(scope="module")
def catalog(
    tmp_path_factory: pytest.TempPathFactory, template_db: Path, codebook_dir: Path, seed_dir: Path
) -> Iterator[Catalog]:
    database = open_copy(tmp_path_factory, template_db, codebook_dir, seed_dir, "export")
    try:
        with database.session() as session:
            ids = build(session)
        yield Catalog(database, ids)
        assert_consistent(database)
    finally:
        database.engine.dispose()


# --------------------------------------------------------------------------------------
# Cell rules (unit level)
# --------------------------------------------------------------------------------------


@pytest.mark.parametrize("prefix", ["=", "+", "-", "@", "\t", "\r", "\n"])
def test_csv_string_cells_are_verbatim(prefix: str) -> None:
    assert csv_cell("string", f"{prefix}SUM(A1)") == f"{prefix}SUM(A1)"


def test_csv_numbers_booleans_dates_and_empty_values_are_unchanged() -> None:
    assert csv_cell("number", -5) == "-5"
    assert csv_cell("number", 2.0) == "2"
    assert csv_cell("number", -2.5) == "-2.5"
    assert csv_cell("bool", True) == "true"
    assert csv_cell("bool", False) == "false"
    assert csv_cell("date", "2012-05") == "2012-05"
    assert csv_cell("string", None) == ""
    assert csv_cell("string", "Plain = text") == "Plain = text"


def test_xlsx_text_replaces_illegal_characters_and_truncates() -> None:
    assert xlsx_text("a\x00b\x0bc\x1fd\te\nf") == "a\ufffdb\ufffdc\ufffdd\te\nf"
    truncated = xlsx_text("y" * 40_000)
    assert len(truncated) == XLSX_MAX_CHARS
    assert truncated.endswith(" [truncated]")
    assert xlsx_text("y" * XLSX_MAX_CHARS) == "y" * XLSX_MAX_CHARS


def test_export_filename_uses_utc() -> None:
    moment = datetime(2026, 10, 28, 17, 5, 9, tzinfo=UTC)
    assert export_filename("project", "csv", moment) == "project-20261028-170509.csv"


# --------------------------------------------------------------------------------------
# CSV
# --------------------------------------------------------------------------------------

DEFAULT_PROJECT_HEADER = [
    "id",
    "slug",
    "name_en",
    "name_uk",
    "in_paper_scope",
    "status",
    "status_label",
    "paper_cluster",
    "paper_cluster_label",
    "ministry_categories",
    "ministry_categories_label",
    "policy_domains",
    "policy_domains_label",
    "wartime_relevance",
    "wartime_relevance_label",
    "derived_launch",
    "derived_launch_start",
    "derived_launch_end",
    "derived_launch_precision",
    "derived_launch_basis",
    "derived_launch_event_id",
    "derived_period",
    "derived_period_label",
    "derived_period_reason",
    "verification_status",
    "verification_status_label",
    "open_comments",
]


def by_header(rows: list[list[str]]) -> list[dict[str, str]]:
    header, *body = rows
    return [dict(zip(header, row, strict=True)) for row in body]


def test_csv_default_columns_and_values(catalog: Catalog) -> None:
    rows = catalog.csv_rows("project", sort="slug")
    assert rows[0] == DEFAULT_PROJECT_HEADER
    records = by_header(rows)
    assert [record["slug"] for record in records] == ["exp-alpha", "exp-beta", "exp-gamma"]
    alpha, beta, gamma = records
    assert alpha["id"] == catalog.ids["X1"]
    assert alpha["name_en"] == '=HYPERLINK("https://example.org")'
    assert alpha["name_uk"] == "Альфа"
    assert alpha["in_paper_scope"] == "true"
    assert (alpha["status"], alpha["status_label"]) == ("active", "Example active")
    assert alpha["policy_domains"] == "domain_b; domain_a"
    assert alpha["policy_domains_label"] == "Example domain_b; Example domain_a"
    assert alpha["derived_launch"] == "2012-05"
    assert (alpha["derived_launch_start"], alpha["derived_launch_end"]) == (
        "2012-05-01",
        "2012-05-31",
    )
    assert alpha["derived_launch_precision"] == "month"
    assert alpha["derived_launch_basis"] == "stated_in_source"
    assert alpha["derived_launch_event_id"] == catalog.ids["E1"]
    assert (alpha["derived_period"], alpha["derived_period_reason"]) == ("pre_invasion", "")
    assert alpha["derived_period_label"] == "Example pre_invasion"
    assert alpha["open_comments"] == "0"
    assert (beta["name_en"], beta["name_uk"]) == ("-5", "@mention")
    assert beta["in_paper_scope"] == "false"
    assert (beta["status"], beta["status_label"]) == ("", "")
    assert (beta["derived_period"], beta["derived_period_reason"]) == ("unknown", "no_date")
    assert gamma["name_en"] == "#N/A"


def test_csv_quotes_every_cell(catalog: Catalog) -> None:
    text = catalog.csv_bytes("project", cols="id,open_comments", sort="slug").decode("utf-8-sig")
    lines = text.splitlines()
    assert lines[0] == '"id","open_comments"'
    assert all(line.startswith('"') and line.endswith('"') for line in lines)


def test_csv_selected_columns_lists_derived_and_system_fields(catalog: Catalog) -> None:
    rows = catalog.csv_rows(
        "project",
        cols="id,aliases,official_urls,existed_by,war_trajectory,created_at,version,comment_count",
        **{"f.slug": "eq:exp-alpha"},
    )
    assert rows[0] == [
        "id",
        "aliases",
        "official_urls",
        "derived_existed_by",
        "derived_existed_by_event_id",
        "derived_war_trajectory",
        "derived_war_trajectory_label",
        "derived_war_trajectory_reason",
        "derived_war_trajectory_event_id",
        "created_at",
        "version",
        "comment_count",
    ]
    record = by_header(rows)[0]
    assert record["aliases"] == "+Alias two; Alias one"
    assert record["official_urls"] == "https://example.org/one; https://example.org/two"
    assert record["derived_war_trajectory"] == "prewar_unchanged"
    assert record["created_at"].endswith("Z")
    assert record["version"] == "1"
    assert record["comment_count"] == "0"


def test_csv_partial_dates_references_numbers_json_and_links(catalog: Catalog) -> None:
    event = by_header(catalog.csv_rows("event", cols="date,subject_id"))[0]
    assert event == {
        "date": "2012-05",
        "date_start": "2012-05-01",
        "date_end": "2012-05-31",
        "date_precision": "month",
        "subject_id": catalog.ids["X1"],
        "subject_id_label": '=HYPERLINK("https://example.org")',
    }
    service = by_header(catalog.csv_rows("service", cols="project_id"))[0]
    assert service["project_id"] == catalog.ids["X1"]
    indicator = by_header(catalog.csv_rows("indicator", cols="value,metric"))[0]
    assert indicator == {"value": "-2.5", "metric": "Budget"}
    term = by_header(
        catalog.csv_rows("vocabulary_term", cols="code,attributes", **{"f.code": "eq:launch"})
    )[0]
    assert term["attributes"] == '{"counts_as_modification":false,"fixed":true}'
    document = by_header(catalog.csv_rows("document", cols="links"))[0]
    assert document["links"] == (
        f'[{{"origin":"manual","state":"confirmed","subject_id":"{catalog.ids["X2"]}"}}]'
    )
    assert document["links_label"] == "-5 (confirmed)"


def test_csv_respects_filters_sort_search_and_deletion_state(catalog: Catalog) -> None:
    def slugs(**params: str) -> list[str]:
        return [record["slug"] for record in by_header(catalog.csv_rows("project", **params))]

    assert slugs(sort="-slug") == ["exp-gamma", "exp-beta", "exp-alpha"]
    assert slugs(**{"f.status": "in:active"}) == ["exp-alpha"]
    assert slugs(q="mention") == ["exp-beta"]
    assert slugs(q="- .") == []
    assert slugs(deleted="only") == ["exp-deleted"]


# --------------------------------------------------------------------------------------
# XLSX
# --------------------------------------------------------------------------------------


def test_xlsx_cells_are_typed_and_never_formulas(catalog: Catalog) -> None:
    workbook = catalog.xlsx("project", sort="slug", cols="id,name_en,in_paper_scope,open_comments")
    sheet = workbook["project"]
    assert workbook.sheetnames == ["project"]
    assert sheet.freeze_panes == "A2"
    header = next(sheet.iter_rows(min_row=1, max_row=1))
    assert [cell.value for cell in header] == ["id", "name_en", "in_paper_scope", "open_comments"]
    assert all(cell.font.b for cell in header)
    rows = list(sheet.iter_rows(min_row=2))
    alpha, beta, gamma = rows
    assert alpha[1].value == '=HYPERLINK("https://example.org")'
    assert alpha[1].data_type == "s"
    assert (beta[1].value, beta[1].data_type) == ("-5", "s")
    assert (gamma[1].value, gamma[1].data_type) == ("#N/A", "s")
    assert (alpha[2].value, alpha[2].data_type) == (True, "b")
    assert (alpha[3].value, alpha[3].data_type) == (0, "n")


def test_xlsx_numbers_empty_cells_and_truncation(catalog: Catalog) -> None:
    indicator = catalog.xlsx("indicator", cols="value,unit")["indicator"]
    value, unit = next(indicator.iter_rows(min_row=2))
    assert (value.value, value.data_type) == (-2.5, "n")
    assert unit.value is None
    terms = catalog.xlsx("vocabulary_term", cols="code,attributes", **{"f.code": "eq:exp_long"})
    code, attributes = next(terms["vocabulary_term"].iter_rows(min_row=2))
    assert code.value == "exp_long"
    assert len(attributes.value) == XLSX_MAX_CHARS
    assert attributes.value.endswith(" [truncated]")
    # CSV keeps the full value
    csv_record = by_header(
        catalog.csv_rows("vocabulary_term", cols="attributes", **{"f.code": "eq:exp_long"})
    )[0]
    assert LONG_NOTE in csv_record["attributes"]


# --------------------------------------------------------------------------------------
# Validation, limits and the export slot
# --------------------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("entity_type", "params"),
    [
        ("project", {"cols": "id,unknown"}),
        ("project", {"cols": "seed_key"}),
        ("document", {"cols": "text"}),
        ("relation", {"cols": "endpoint"}),
        ("project", {"format": "pdf"}),
    ],
)
def test_invalid_export_requests(
    catalog: Catalog, entity_type: str, params: dict[str, str]
) -> None:
    with catalog.db.session() as session, pytest.raises(ValidationError):
        start_export(session, entity_type, params)


def test_the_meta_flag_exportable_names_exactly_the_exportable_columns(catalog: Catalog) -> None:
    # The UI offers as table and export columns only the fields that /api/meta marks
    # ``exportable`` (UI §3); an export of any other column is refused.
    for spec in ENTITY_TYPES.values():
        if spec.name == "comment":
            continue
        for field in spec.fields:
            if field.internal:
                continue
            with catalog.db.session() as session:
                if field.meta()["exportable"]:
                    start_export(session, spec.name, {"cols": f"id,{field.name}"}).close()
                else:
                    with pytest.raises(ValidationError):
                        start_export(session, spec.name, {"cols": f"id,{field.name}"})
    assert not ENTITY_TYPES["document"].field("text").exportable
    assert not ENTITY_TYPES["relation"].field("endpoint").exportable
    assert ENTITY_TYPES["project"].field("open_comments").exportable


def test_unknown_types_cannot_be_exported(catalog: Catalog) -> None:
    with catalog.db.session() as session, pytest.raises(NotFoundError):
        start_export(session, "comment", {})


def test_export_jobs(catalog: Catalog) -> None:
    with catalog.db.session() as session:
        first = start_export(session, "project", {})
        assert first.media_type == "text/csv; charset=utf-8"
        assert first.filename.startswith("project-") and first.filename.endswith(".csv")
        # Exports do not wait for each other.
        second = start_export(session, "service", {})
        b"".join(second.iter_csv())
        with pytest.raises(RuntimeError):
            list(second.iter_csv())  # the output is produced once
        first.close()
        first.close()  # idempotent
        third = start_export(session, "service", {"format": "xlsx"})
        assert third.media_type.endswith("spreadsheetml.sheet")
        assert third.filename.endswith(".xlsx")
        third.write_xlsx(io.BytesIO())


# --------------------------------------------------------------------------------------
# Cost: exports load only what the selected columns need
# --------------------------------------------------------------------------------------


def test_exports_load_only_what_the_columns_need(
    catalog: Catalog, monkeypatch: pytest.MonkeyPatch
) -> None:
    # Exports used to serialise every row as a full API item (links, virtual counts,
    # labels and a deep copy of field_meta) whatever ``cols`` asked for.
    def not_needed(*_args: Any, **_kwargs: Any) -> Any:
        raise AssertionError("loaded although no selected column reads it")

    for name in ("load_links", "derived_rows", "virtual_values", "ref_infos"):
        monkeypatch.setattr(export, name, not_needed)
    with count_statements(catalog.db.engine) as statements:
        rows = catalog.csv_rows("project", cols="id,slug,name_en,status,created_by")
    assert rows[0] == ["id", "slug", "name_en", "status", "status_label", "created_by"]
    assert any(row[1] == "exp-alpha" for row in rows[1:])
    assert not any("field_meta" in statement for statement in statements)


def test_export_needs_follow_the_columns() -> None:
    spec = ENTITY_TYPES["document"]
    assert export.export_needs(spec, ["id"]) == export.ExportNeeds()
    needs = export.export_needs(
        spec, ["id", "title", "published", "links", "publisher_org_id", "open_comments", "version"]
    )
    assert needs.columns == (
        "title",
        "published_start",
        "published_end",
        "published_precision",
        "publisher_org_id",
        "version",
    )
    assert [field.name for field in needs.fields] == [
        "title",
        "published",
        "links",
        "publisher_org_id",
    ]
    assert (needs.virtual, needs.system, needs.ref_fields) == (
        ("open_comments",),
        ("version",),
        ("publisher_org_id",),
    )
    assert needs.links and not needs.derived
    assert export.export_needs(ENTITY_TYPES["project"], ["launch"]).derived


def test_export_queries_do_not_grow_with_the_rows(db: Db) -> None:
    def statement_count(**params: str) -> int:
        with db.session() as session, count_statements(db.engine) as statements:
            job = start_export(session, "document", {"format": "csv", **params})
            b"".join(job.iter_csv())
        return len(statements)

    columns = "id,title,published,links,publisher_org_id,open_comments,linked_in_paper_scope"
    with db.session() as session:
        project_id = create_project(session, name_en="Export scaling project").entity_id
        for index in range(2):
            create_document(
                session,
                f"Scaling {index}",
                links=[{"subject_id": project_id, "state": "confirmed"}],
            )
    few = statement_count(cols=columns)
    with db.session() as session:
        for index in range(2, 30):
            create_document(
                session,
                f"Scaling {index}",
                links=[{"subject_id": project_id, "state": "confirmed"}],
            )
    assert statement_count(cols=columns) == few
