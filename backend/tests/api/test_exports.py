"""Exports over HTTP (API §5): status, headers and content."""

from __future__ import annotations

import csv
import io
import re

from openpyxl import load_workbook

from mdt.domain.export import start_export
from tests.api.helpers import ApiClient, error_of
from tests.conftest import Db

EXPORT = "/api/entities/project/export"


def csv_rows(content: bytes) -> list[list[str]]:
    assert content.startswith("﻿".encode())
    return list(csv.reader(io.StringIO(content.decode("utf-8-sig"))))


def test_csv_export_of_all_projects(client: ApiClient) -> None:
    response = client.get(EXPORT, {"format": "csv"})
    assert response.status_code == 200
    assert response.headers["content-type"] == "text/csv; charset=utf-8"
    assert response.headers["cache-control"] == "no-store"
    assert re.fullmatch(
        r'attachment; filename="project-[0-9]{8}-[0-9]{6}\.csv"',
        response.headers["content-disposition"],
    )
    rows = csv_rows(response.content)
    header, data = rows[0], rows[1:]
    assert header[0] == "id" and "slug" in header and "paper_cluster_label" in header
    assert len(data) == 20
    assert all(row[0].startswith("prj_") for row in data)


def test_csv_export_follows_filters_sort_and_columns(client: ApiClient) -> None:
    response = client.get(
        EXPORT, {"f.in_paper_scope": "eq:true", "sort": "-slug", "cols": "slug,in_paper_scope"}
    )
    rows = csv_rows(response.content)
    assert rows[0] == ["slug", "in_paper_scope"]
    slugs = [row[0] for row in rows[1:]]
    assert len(slugs) == 19 and slugs == sorted(slugs, reverse=True)
    assert {row[1] for row in rows[1:]} == {"true"}


def test_csv_cells_hold_the_values_verbatim(client: ApiClient) -> None:
    client.create_project(name_en="=1+1", notes="@SUM(A1)")
    response = client.get(EXPORT, {"cols": "name_en,notes", "f.name_en": "eq:=1+1"})
    assert csv_rows(response.content)[1] == ["=1+1", "@SUM(A1)"]


def test_xlsx_export(client: ApiClient) -> None:
    client.create_project(name_en='=HYPERLINK("https://example.org")')
    response = client.get(EXPORT, {"format": "xlsx", "cols": "slug,name_en,in_paper_scope"})
    assert response.status_code == 200
    assert response.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert response.headers["content-length"] == str(len(response.content))
    assert response.headers["content-disposition"].endswith('.xlsx"')
    assert response.headers["cache-control"] == "no-store"
    workbook = load_workbook(io.BytesIO(response.content))
    sheet = workbook["project"]
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0] == ("slug", "name_en", "in_paper_scope")
    assert len(rows) == 22
    formula_cells = [row for row in sheet.iter_rows(min_row=2) if str(row[1].value).startswith("=")]
    assert len(formula_cells) == 1 and formula_cells[0][1].data_type == "s"
    assert {row[2] for row in rows[1:]} == {True, False}


def test_export_parameters_are_validated(client: ApiClient) -> None:
    cases = [
        ({"format": "pdf"}, "validation_error"),
        ({"cols": "slug,nope"}, "validation_error"),
        ({"f.slug": "between:a"}, "bad_filter"),
        ({"sort": "aliases"}, "bad_sort"),
    ]
    for params, code in cases:
        response = client.get(EXPORT, params)
        assert (response.status_code, error_of(response)["code"]) == (400, code), params
    assert client.get("/api/entities/comment/export").status_code == 404


def test_exports_do_not_wait_for_each_other(client: ApiClient, api_db: Db) -> None:
    with api_db.session() as session:
        job = start_export(session, "project", {})
        try:
            assert client.get(EXPORT).status_code == 200
            assert client.get(EXPORT, {"format": "xlsx"}).status_code == 200
        finally:
            job.close()
