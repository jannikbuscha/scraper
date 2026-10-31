"""Reading the uploaded tables of an import (docs/COLLECTORS.md §9.2).

* **CSV:** UTF-8 with or without BOM (anything else is refused: "save the file as UTF-8");
  the delimiter is the one of ``,``, ``;`` and tab that occurs most often in the header
  line; quoting per RFC 4180.
* **XLSX:** openpyxl in read-only mode with cached formula values; the chosen sheet
  (default: the first). Numbers become their decimal text (integral values without
  ``.0``), date cells ``YYYY-MM-DD``.
* **Header:** the first row; an empty header cell is named ``column_<n>``; duplicate
  headers are refused. Empty rows are skipped; rows keep their numbers (header = row 1).
"""

from __future__ import annotations

import csv
import io
import zipfile
from collections.abc import Sequence
from dataclasses import dataclass
from datetime import date, datetime, time
from decimal import Decimal
from typing import Any, Literal

import openpyxl
from openpyxl.utils.exceptions import InvalidFileException
from sqlalchemy.orm import Session

from mdt.collectors.sources import MANUAL
from mdt.collectors.store import XLSX, SnapshotRecord, SnapshotStore
from mdt.config import Settings
from mdt.domain.errors import ValidationError

CSV_DELIMITERS = (",", ";", "\t")


@dataclass(frozen=True)
class Table:
    """The cells of one sheet as text: the column names and the non-empty data rows with
    their row numbers (the header is row 1)."""

    columns: tuple[str, ...]
    rows: tuple[tuple[int, tuple[str, ...]], ...]
    sheets: tuple[str, ...] | None = None
    sheet: str | None = None


def _file_error(message: str) -> ValidationError:
    return ValidationError.single("upload_id", message)


def column_names(header: Sequence[str]) -> tuple[str, ...]:
    """Column names from the header row: trimmed; an empty cell is ``column_<n>``;
    duplicates are refused."""
    names = tuple(cell.strip() or f"column_{index}" for index, cell in enumerate(header, start=1))
    duplicates = sorted({name for name in names if names.count(name) > 1})
    if duplicates:
        raise _file_error(f"duplicate column headers: {', '.join(duplicates)}")
    return names


def _table(rows: Sequence[tuple[int, Sequence[str]]], **sheet: Any) -> Table:
    if not rows or not any(cell.strip() for cell in rows[0][1]):
        raise _file_error("the file has no header row")
    columns = column_names(rows[0][1])
    data = tuple(
        (number, tuple(cells)) for number, cells in rows[1:] if any(cell.strip() for cell in cells)
    )
    return Table(columns, data, **sheet)


def detect_delimiter(header_line: str) -> str:
    """The one of ``,``, ``;`` and tab that occurs most often in the header line."""
    counts = [header_line.count(delimiter) for delimiter in CSV_DELIMITERS]
    return CSV_DELIMITERS[counts.index(max(counts))]


def read_csv(body: bytes) -> Table:
    """A UTF-8 CSV file (with or without BOM), RFC 4180 quoting."""
    try:
        text = body.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise _file_error("the CSV file is not UTF-8; save the file as UTF-8") from exc
    delimiter = detect_delimiter(text.split("\n", 1)[0])
    try:
        records = list(csv.reader(io.StringIO(text, newline=""), delimiter=delimiter))
    except csv.Error as exc:
        raise _file_error("the CSV file cannot be read") from exc
    return _table([(number, record) for number, record in enumerate(records, start=1)])


def cell_text(value: Any) -> str:
    """An XLSX cell as text: numbers as decimal text (integral values without ``.0``),
    date cells as ``YYYY-MM-DD`` (a time of day is kept), booleans as ``true``/``false``."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return format(Decimal(repr(value)), "f")
    if isinstance(value, datetime):
        if value.time() == time(0, 0):
            return value.date().isoformat()
        return value.isoformat()
    if isinstance(value, date | time):
        return value.isoformat()
    return str(value)


def read_xlsx(body: bytes, sheet: str | None) -> Table:
    """One sheet of an XLSX workbook (default: the first), with cached formula values."""
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(body), read_only=True, data_only=True)
    except (InvalidFileException, zipfile.BadZipFile, KeyError, OSError, ValueError) as exc:
        raise _file_error("the XLSX file cannot be read") from exc
    try:
        names = tuple(workbook.sheetnames)
        chosen = sheet if sheet is not None else (names[0] if names else None)
        if chosen is None or chosen not in names:
            raise ValidationError.single("sheet", "the workbook has no sheet of this name")
        worksheet = workbook[chosen]
        start = worksheet.min_row or 1
        rows = [
            (number, [cell_text(value) for value in values])
            for number, values in enumerate(worksheet.iter_rows(values_only=True), start=start)
        ]
    finally:
        workbook.close()
    return _table(rows, sheets=names, sheet=chosen)


def file_kind(record: SnapshotRecord) -> Literal["csv", "xlsx"]:
    """``csv`` or ``xlsx`` for an upload; anything else is refused."""
    if record.content_type == XLSX:
        return "xlsx"
    if record.content_type in ("text/csv", "text/plain"):
        return "csv"
    if (record.filename or "").lower().endswith((".csv", ".tsv")):
        raise _file_error("the CSV file is not UTF-8; save the file as UTF-8")
    raise _file_error("upload a CSV or XLSX file")


def read_table(record: SnapshotRecord, body: bytes, sheet: str | None) -> Table:
    """The table of an uploaded CSV or XLSX file (``sheet`` only for XLSX)."""
    if file_kind(record) == "xlsx":
        return read_xlsx(body, sheet)
    if sheet is not None:
        raise ValidationError.single("sheet", "a CSV file has no sheets")
    return read_csv(body)


def load_upload(
    session: Session, settings: Settings, upload_id: Any
) -> tuple[SnapshotRecord, bytes]:
    """An upload snapshot (``via = upload``) and its body."""
    if not isinstance(upload_id, str) or not upload_id:
        raise ValidationError.single("upload_id", "is required")
    store = SnapshotStore(settings.raw_path)
    record = store.record(session, upload_id)
    if record is None or record.via != "upload" or record.source_id != MANUAL:
        raise ValidationError.single("upload_id", "names no uploaded file")
    return record, store.read_body(record.sha256)
